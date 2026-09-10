#!/usr/bin/env python3
"""
把 HTML 幻灯片导出成 pptx（截图版）。

每页截成高清图铺满 16:9 版面，视觉 100% 保真——字体、色块、间距和浏览器里
一模一样，不会因为目标机器缺字体而跑版。代价是文字不可编辑。

需要可编辑版本时，不要指望自动转换：那要按内容用 python-pptx 逐页重建文本框
和形状，属于一次性手工活，视觉也一定有偏差（换机器字体会被替换）。取舍见
SKILL.md「导出 pptx」一节。

用法：
    python3 export_pptx.py deck.html
    python3 export_pptx.py deck.html --out 分享.pptx --scale 3

依赖：playwright、python-pptx。浏览器复用本机缓存，不联网下载。
"""
import argparse, os, sys, glob, tempfile, shutil
from pathlib import Path


def activate_slide(page, index):
    page.evaluate(
        """index => {
          const slides = [...document.querySelectorAll('.slide, [data-slide], section.slide')];
          const dots = [...document.querySelectorAll('.progress button')];
          if (dots[index]) { dots[index].click(); return; }
          slides.forEach((slide, current) => slide.classList.toggle('active', current === index));
        }""",
        index,
    )


def find_browser():
    """复用本机 playwright 缓存，绝不触发下载。"""
    root = os.path.expanduser("~/Library/Caches/ms-playwright")
    if not os.path.isdir(root):
        root = os.path.expanduser("~/.cache/ms-playwright")
    pats = ["chromium_headless_shell-*/*/chrome-headless-shell",
            "chromium-*/chrome-*/Chromium",
            "chromium-*/chrome-linux/chrome"]
    hits = []
    for p in pats:
        hits += glob.glob(os.path.join(root, p))
    return sorted(hits)[-1] if hits else None


def shoot(html, tmp, scale):
    from playwright.sync_api import sync_playwright
    f = Path(html).expanduser()
    if not f.is_file():
        print(f"文件不存在：{html}", file=sys.stderr); sys.exit(2)
    exe = find_browser()
    if not exe:
        print("找不到本机 playwright 浏览器缓存，不要联网重装。", file=sys.stderr); sys.exit(2)

    titles = []
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=exe)
        # 视口取宽屏：deck 容器通常有 max-width，再大也不会更清晰，靠 scale 提分辨率
        pg = b.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=scale)
        pg.goto(f.resolve().as_uri()); pg.wait_for_timeout(1200)
        stage = pg.query_selector("#stage, .stage, main")
        n = pg.evaluate("document.querySelectorAll('.slide, [data-slide], section.slide').length")
        if not n:
            print("没找到幻灯片节点（.slide / [data-slide] / section.slide）", file=sys.stderr)
            sys.exit(2)
        for i in range(n):
            activate_slide(pg, i)
            pg.wait_for_timeout(650)   # 等翻页动画落定，短了会截到过渡帧
            (stage or pg).screenshot(path=os.path.join(tmp, f"s{i+1}.png"))
            titles.append(pg.evaluate(
                "(document.querySelector('.slide.active, section.active')||{}).dataset?.title || ''"))
        box = stage.bounding_box() if stage else None
        if box:
            print(f"截图 {n} 页，每页 {int(box['width']*scale)}x{int(box['height']*scale)} px")
        b.close()
    return titles


def build(tmp, titles, out):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9
    blank = prs.slide_layouts[6]
    for i, t in enumerate(titles, 1):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(os.path.join(tmp, f"s{i}.png"), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)
        # 备注写页码和标题，放映时演讲者视图能看到，便于定位
        s.notes_slide.notes_text_frame.text = f"P{i}" + (f" · {t}" if t else "")
    prs.save(out)


def main():
    ap = argparse.ArgumentParser(description="HTML 幻灯片导出 pptx（截图版）")
    ap.add_argument("html")
    ap.add_argument("--out", default="", help="输出路径，默认与 html 同名同目录")
    ap.add_argument("--scale", type=int, default=3, help="截图倍率，3 足够 4K 投影")
    a = ap.parse_args()

    out = a.out or str(Path(a.html).with_suffix(".pptx"))
    tmp = tempfile.mkdtemp(prefix="deck-")
    try:
        titles = shoot(a.html, tmp, a.scale)
        build(tmp, titles, out)
        print(f"已生成：{out}（{os.path.getsize(out)/1024/1024:.1f} MB，{len(titles)} 页）")
        print("提示：文字不可编辑。现场优先级 HTML 全屏 > 截图 pptx > 可编辑 pptx。")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
